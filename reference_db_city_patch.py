"""
reference_db_city_patch.py
==========================
Copy-paste these functions INTO your existing reference_db.py file.

These replace / add:
  - get_city_sheet_names()
  - load_city_db_sheet(sheet_names_str)
  - load_city_reference(n_cities)
  - fallback_city_sheet_names_from_workbook(path)   ← kept for safety

They read from the  city_reference  PostgreSQL table populated by
migrate_city_db.py.
"""

from __future__ import annotations

import os
import re
from functools import lru_cache

# ── your existing _get_conn() / engine is already in reference_db.py ────────
# These functions reuse it.  Import or reference as needed.


# ---------------------------------------------------------------------------
# City sheet names
# ---------------------------------------------------------------------------

def get_city_sheet_names() -> list[str]:
    """
    Return the distinct sheet_name values from city_reference,
    ordered by the total potential_impressions descending so the
    most-used sheets appear first.
    """
    try:
        conn = _get_conn()          # your existing helper
        cur  = conn.cursor()
        cur.execute("""
            SELECT sheet_name
            FROM   city_reference
            GROUP  BY sheet_name
            ORDER  BY SUM(potential_impressions) DESC NULLS LAST
        """)
        names = [row[0] for row in cur.fetchall()]
        cur.close()
        conn.close()
        return names
    except Exception as e:
        print(f"[DB] get_city_sheet_names failed: {e}")
        return []


def fallback_city_sheet_names_from_workbook(path: str) -> list[str]:
    """Fallback: read sheet names directly from the Excel file."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        names = wb.sheetnames
        wb.close()
        return names
    except Exception:
        return []


# ---------------------------------------------------------------------------
# Load city data for report generation
# ---------------------------------------------------------------------------

def load_city_db_sheet(sheet_names_str: str) -> list[dict]:
    """
    Load city rows for one or more sheet names (comma-separated).

    Returns list of dicts:
        [{"name": str, "weight": float, "creatives": []}, ...]

    weight = potential_impressions  (used for impression distribution)
    creatives is left empty — App.py fills it from File 1 line items.
    """
    if not sheet_names_str or not sheet_names_str.strip():
        return []

    sheet_list = [s.strip() for s in sheet_names_str.split(",") if s.strip()]
    if not sheet_list:
        return []

    try:
        conn = _get_conn()
        cur  = conn.cursor()

        placeholders = ",".join(["%s"] * len(sheet_list))
        cur.execute(
            f"""
            SELECT city_name,
                   COALESCE(potential_impressions, unique_cookies, 1) AS weight
            FROM   city_reference
            WHERE  sheet_name IN ({placeholders})
            ORDER  BY weight DESC NULLS LAST
            """,
            sheet_list,
        )
        rows = cur.fetchall()
        cur.close()
        conn.close()

        seen: set[str] = set()
        results: list[dict] = []
        for city_name, weight in rows:
            key = city_name.strip().lower()
            if key in seen:
                continue
            seen.add(key)
            results.append({
                "name":      city_name.strip(),
                "weight":    float(weight) if weight else 1.0,
                "creatives": [],
            })
        return results

    except Exception as e:
        print(f"[DB] load_city_db_sheet failed: {e}")
        return []


# ---------------------------------------------------------------------------
# Lightweight city reference (top-N cities for fallback weighting)
# ---------------------------------------------------------------------------

def load_city_reference(n_cities: int = 50) -> list[tuple[str, float]]:
    """
    Return top-N (city_name, weight) tuples from the Australia master sheet,
    ordered by potential_impressions descending.
    Used as a fallback weight source in build_sheet8_city().
    """
    try:
        conn = _get_conn()
        cur  = conn.cursor()
        cur.execute(
            """
            SELECT city_name,
                   COALESCE(potential_impressions, unique_cookies, 1) AS weight
            FROM   city_reference
            WHERE  sheet_name = 'Australia'
            ORDER  BY weight DESC NULLS LAST
            LIMIT  %s
            """,
            (n_cities,),
        )
        rows = cur.fetchall()
        cur.close()
        conn.close()
        return [(r[0], float(r[1])) for r in rows]
    except Exception as e:
        print(f"[DB] load_city_reference failed: {e}")
        return []
